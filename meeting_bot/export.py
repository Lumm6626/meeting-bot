# -*- coding: utf-8 -*-
"""导出：逐字稿 Markdown + 纪要 Word。"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document

from .config import Settings
from .transcript import TranscriptStore
from .translator import MinutesResult, TranslationError, Translator

MINUTES_FAILED_SUMMARY = "（纪要生成失败，请查看逐字稿）"
NO_TRANSLATOR_SUMMARY = "（翻译接口不可用，未生成纪要）"


def _safe_title(title: str) -> str:
    r"""净化标题：替换路径分隔符、Windows 非法字符与控制符，防止路径穿越。

    - `..` / `../` 中的斜杠与点被处理，标题无法逃出 output_dir；
    - `\ / : * ? " < > |` 与全部 C0 控制符（含 \t \r \n）及 DEL（\x7f）替换为 `_`，
      避免 Windows 文件名 OSError；
    - 净化后为空（空串、纯点、纯空格；控制符已先变成 `_`）回退「未命名会议」，最长 80 字符；
    - strip 在截断后再兜一次，避免超长标题截断时重新露出尾随空格/点。
    """
    # \x5c 即反斜杠（写成转义码，避免被工具链/编辑器吞掉一层）
    t = re.sub(r'[\x00-\x1f\x7f\x5c/:*?"<>|]', "_", title).strip(" .")
    return t[:80].strip(" .") or "未命名会议"


_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _xml_safe(s: str) -> str:
    r"""删除 XML 1.0 非法控制符（保留 \t \n \r），避免 python-docx 抛 ValueError。

    python-docx 将字符串写入 XML，遇到 0x00-0x08 / 0x0b / 0x0c / 0x0e-0x1f 会报
    "All strings must be XML compatible"；LLM 生成的纪要内容不可信，统一先清洗。
    """
    return _XML_ILLEGAL.sub("", s)


def meeting_dir(settings: Settings, title: str, started_at: datetime) -> Path:
    """产物目录：标题已含日期时直接用标题（避免「日期_时间_日期 时间」重复）。"""
    safe = _safe_title(title)
    if re.match(r"^\d{4}-\d{2}-\d{2}", safe):
        d = settings.output_dir / safe
    else:
        d = settings.output_dir / f"{started_at:%Y-%m-%d_%H%M}_{safe}"
    # 防线：净化后仍校验落点包含关系，且先校验后建目录（越界不留半成品）
    assert d.resolve().is_relative_to(settings.output_dir.resolve())
    d.mkdir(parents=True, exist_ok=True)
    return d


def write_transcript(store: TranscriptStore, directory: Path,
                     title: str, started_at: datetime,
                     md_text: str | None = None) -> Path:
    """写逐字稿；md_text 非空时复用调用方缓存的 Markdown，避免重复渲染。"""
    path = directory / "逐字稿.md"
    path.write_text(md_text if md_text is not None else store.to_markdown(title, started_at),
                    encoding="utf-8")
    return path


def minutes_to_docx(minutes: MinutesResult, path: Path) -> Path:
    doc = Document()
    doc.add_heading(_xml_safe(minutes.title), level=1)
    doc.add_paragraph(_xml_safe(minutes.summary))
    doc.add_heading("讨论要点", level=2)
    for p in minutes.key_points:
        doc.add_paragraph(_xml_safe(p), style="List Bullet")
    doc.add_heading("会议结论", level=2)
    for p in minutes.decisions:
        doc.add_paragraph(_xml_safe(p), style="List Bullet")
    doc.add_heading("行动项", level=2)
    for a in minutes.action_items:
        task = _xml_safe((a.task or "").strip()) or "（未命名任务）"
        owner = _xml_safe((a.owner or "").strip()) or "待定"
        due = f"（截止：{_xml_safe(a.due)}）" if a.due else ""
        doc.add_paragraph(f"{task} — 负责人：{owner}{due}", style="List Bullet")
    doc.add_heading("待跟进问题", level=2)
    for p in minutes.open_questions:
        doc.add_paragraph(_xml_safe(p), style="List Bullet")
    doc.save(str(path))
    return path


def export_all(store: TranscriptStore, settings: Settings, title: str,
               started_at: datetime,
               translator: Translator | None = None) -> tuple[Path, Path]:
    directory = meeting_dir(settings, title, started_at)
    md = store.to_markdown(title, started_at)
    transcript_path = write_transcript(store, directory, title, started_at, md_text=md)
    if translator is not None:
        try:
            minutes = translator.minutes(md)
        except TranslationError:
            # 纪要失败不阻断导出：逐字稿已落盘，纪要降级为占位内容
            minutes = MinutesResult(title=_safe_title(title), summary=MINUTES_FAILED_SUMMARY)
    else:
        minutes = MinutesResult(title=_safe_title(title), summary=NO_TRANSLATOR_SUMMARY)
    minutes_path = minutes_to_docx(minutes, directory / "纪要.docx")
    return transcript_path, minutes_path
